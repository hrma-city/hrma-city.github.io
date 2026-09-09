# -*- coding: utf-8 -*-
"""
HRMA · Excel 实操模板包 —— 公共样式与辅助函数
所有配色取自 SPEC 7.1；所有术语/阈值取自 SPEC 三、四章。
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# ---------------------------------------------------------------- 配色（SPEC 7.1）
NAVY = "1B3A5C"
NAVY2 = "2C5282"
GOLD = "B8894A"
GOLD_SOFT = "F5EDE0"
TEAL = "2E7D74"
RED = "C0392B"        # 涨 / 警告（中国习惯）
GREEN = "2F7A3E"      # 跌 / 下降
AMBER = "D98324"
INK = "1F2933"
INK2 = "52606D"
LINE = "E1E5EA"
BG = "FAFAF8"
BG2 = "F2F0EB"
WHITE = "FFFFFF"

INPUT_BG = "FFF7D6"   # 输入区：浅黄
REF_BG = "EDEDED"     # 参考区：灰
RED_LIGHT = "FBE3E0"
GREEN_LIGHT = "E2F0E4"
AMBER_LIGHT = "FCEFD9"
TEAL_LIGHT = "DFF0ED"
GREY_LIGHT = "EDEDED"
GOLD_LIGHT = "F7EFE1"

FONT = "PingFang SC"

# ---------------------------------------------------------------- 数字格式
MONEY = "#,##0"
MONEY2 = "#,##0.00"
PCT1 = "0.0%"
PCT2 = "0.00%"
NUM = "#,##0"
NUM1 = "#,##0.0"
IDX = "0.0"
DATE = "yyyy-mm-dd"
DT = "yyyy-mm-dd hh:mm"
INT = "0"

# ---------------------------------------------------------------- 边框
_s = Side(style="thin", color=LINE)
BORDER = Border(left=_s, right=_s, top=_s, bottom=_s)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def box(horizontal="center", vertical="center", wrap=False):
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


# ---------------------------------------------------------------- 单元格写入
KIND_STYLE = {
    # kind: (fill, font_color, bold, italic)
    "h": (NAVY, WHITE, True, False),    # 标题行：深蓝底白字
    "i": (INPUT_BG, INK, False, False),  # 输入：浅黄
    "f": (WHITE, INK, False, False),     # 公式：白底
    "r": (REF_BG, INK2, False, True),    # 参考：灰底
    "l": (WHITE, INK, False, False),     # 标签
    "s": (GOLD_SOFT, NAVY, True, False),  # 分区标题
    "t": (WHITE, INK, True, False),      # 小标题
    "n": (WHITE, INK2, False, True),     # 注释
}


def put(ws, addr, value, kind="l", fmt=None, align="center", wrap=False,
        bold=None, italic=None, border=True):
    """写入一个单元格。kind 决定底色与字色。"""
    c = ws[addr]
    c.value = value
    bg, fc, bd, it = KIND_STYLE[kind]
    c.fill = fill(bg)
    if bold is None:
        bold = bd
    if italic is None:
        italic = it
    c.font = Font(name=FONT, size=10, bold=bold, color=fc, italic=italic)
    c.alignment = box(align, "center", wrap)
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BORDER
    return c


def row(ws, r, start_col, cells, border=True):
    """cells: 列表，元素为 value 或 (value, kind, fmt) 或 (value, kind, fmt, align)。"""
    for i, spec in enumerate(cells):
        col = start_col + i
        if not isinstance(spec, (list, tuple)):
            spec = (spec, "l", None)
        v = spec[0]
        kind = spec[1] if len(spec) > 1 else "l"
        fmt = spec[2] if len(spec) > 2 else None
        align = spec[3] if len(spec) > 3 else "center"
        if v is not None and v != "":
            put(ws, f"{get_column_letter(col)}{r}", v, kind, fmt, align, border=border)
        else:
            put(ws, f"{get_column_letter(col)}{r}", None, kind, fmt, align, border=border)
    return r + 1


def header_row(ws, r, start_col, headers, height=32, widths=None):
    for i, h in enumerate(headers):
        put(ws, f"{get_column_letter(start_col + i)}{r}", h, "h", None, "center", wrap=True)
    ws.row_dimensions[r].height = height
    if widths:
        set_widths(ws, widths)
    return r + 1


def set_widths(ws, widths, start_col=1):
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + i)].width = w


def title(ws, text, lede, last_col="F", height=30):
    put(ws, "A1", text, "h", None, "left", border=False)
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=WHITE)
    ws["A1"].fill = fill(NAVY)
    for col in range(2, column_index(last_col) + 1):
        ws.cell(row=1, column=col).fill = fill(NAVY)
    ws.merge_cells(f"A1:{last_col}1")
    ws.row_dimensions[1].height = height

    put(ws, "A2", lede, "n", None, "left", wrap=True, border=False)
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color=INK2)
    ws.merge_cells(f"A2:{last_col}2")
    ws.row_dimensions[2].height = 26
    return 3


def section(ws, r, text, last_col="F"):
    put(ws, f"A{r}", text, "s", None, "left", border=False)
    ws[f"A{r}"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    for col in range(2, column_index(last_col) + 1):
        ws.cell(row=r, column=col).fill = fill(GOLD_SOFT)
    ws.merge_cells(f"A{r}:{last_col}{r}")
    ws.row_dimensions[r].height = 22
    return r + 1


def column_index(letter):
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n


def para(ws, r, text, last_col="F", height=None, kind="l"):
    ws.merge_cells(f"A{r}:{last_col}{r}")
    put(ws, f"A{r}", text, kind, None, "left", wrap=True, border=False)
    if height:
        ws.row_dimensions[r].height = height
    return r + 1


# ---------------------------------------------------------------- 数据验证
def dv_list(ws, rng, items, prompt=None):
    s = ",".join(items)
    d = DataValidation(type="list", formula1=f'"{s}"', allow_blank=True, showErrorMessage=True)
    if prompt:
        d.prompt = prompt
        d.promptTitle = "请选择"
        d.showInputMessage = True
    ws.add_data_validation(d)
    d.add(rng)
    return d


def dv_num(ws, rng, lo, hi, kind="decimal"):
    d = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi),
                       allow_blank=True, showErrorMessage=True)
    d.errorTitle = "超出范围"
    d.error = f"请输入 {lo} ~ {hi} 之间的数值"
    ws.add_data_validation(d)
    d.add(rng)
    return d


def dv_date(ws, rng):
    d = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)",
                       allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)
    return d


# ---------------------------------------------------------------- 条件格式
def cf_updown(ws, rng):
    """涨红 / 跌绿 / 持平灰（SPEC 7.1 数字涨跌配色）。"""
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="greaterThan", formula=["0"],
        fill=fill(RED_LIGHT), font=Font(name=FONT, size=10, color=RED, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="lessThan", formula=["0"],
        fill=fill(GREEN_LIGHT), font=Font(name=FONT, size=10, color=GREEN, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=["0"],
        fill=fill(GREY_LIGHT), font=Font(name=FONT, size=10, color=INK2)))


def cf_scale(ws, rng, good_lo, good_hi, warn_lo, warn_hi,
             good=TEAL, warn=AMBER, bad=RED,
             good_bg=TEAL_LIGHT, warn_bg=AMBER_LIGHT, bad_bg=RED_LIGHT):
    """区间判定配色：达标区间=青绿，警戒区间=琥珀，其余=红。"""
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="between", formula=[str(good_lo), str(good_hi)],
        fill=fill(good_bg), font=Font(name=FONT, size=10, color=good, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="between", formula=[str(warn_lo), str(warn_hi)],
        fill=fill(warn_bg), font=Font(name=FONT, size=10, color=warn, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="lessThan", formula=[str(good_lo)],
        fill=fill(bad_bg), font=Font(name=FONT, size=10, color=bad, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="greaterThan", formula=[str(good_hi)],
        fill=fill(bad_bg), font=Font(name=FONT, size=10, color=bad, bold=True)))


def cf_rule(ws, rng, formula_expr, color=RED, bg=None, bold=True):
    """自定义公式条件格式（公式相对 range 左上角）。"""
    f = Font(name=FONT, size=10, color=color, bold=bold)
    kw = dict(font=f)
    if bg:
        kw["fill"] = fill(bg)
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula_expr], **kw))


# ---------------------------------------------------------------- 说明页
COLOR_RULES = [
    ("输入区", "浅黄底 + 边框", INPUT_BG, "由学员填写，例如日期、房晚、价格"),
    ("公式区", "白底 + 边框", WHITE, "已预置 Excel 公式，改动会导致结果错误，请勿手改"),
    ("参考区", "灰底 + 斜体", REF_BG, "SPEC 固定参数/判定表，一般不需要改动"),
    ("标题行", "深蓝底白字", NAVY, "字段名，冻结在首行"),
]


def write_intro(ws, book_no, book_name, purpose, dict_rows, steps, codex,
                lede="", extra_sections=None, last_col="E"):
    """生成统一的『说明』页。每个工作簿第一张 sheet。"""
    ws.sheet_view.showGridLines = False
    set_widths(ws, [8, 24, 14, 34, 46])
    r = title(ws, f"{book_no} · {book_name} ｜ 使用说明",
              lede or "填黄色格子，读白色格子的结果。公式已全部预置，不需要自己写任何公式。",
              last_col=last_col)

    r = section(ws, r, "一、用途", last_col)
    for line in purpose:
        r = para(ws, r, line, last_col, height=18)
    r += 1

    r = section(ws, r, "二、填色规则", last_col)
    for name, style_desc, color, desc in COLOR_RULES:
        put(ws, f"A{r}", "", "l", None, "center", border=False)
        ws[f"A{r}"].fill = fill(color)
        ws[f"A{r}"].border = BORDER
        put(ws, f"B{r}", name, "l", None, "left")
        put(ws, f"C{r}", style_desc, "r", None, "left")
        ws.merge_cells(f"D{r}:{last_col}{r}")
        put(ws, f"D{r}", desc, "l", None, "left", wrap=True)
        ws.row_dimensions[r].height = 20
        r += 1
    r += 1

    r = section(ws, r, "三、字段字典", last_col)
    r = header_row(ws, r, 1, ["列", "字段名称", "类型", "口径 / 预置公式", "说明"])
    for d in dict_rows:
        code, name, typ, formula_txt, note = d
        typ_kind = {"输入": "i", "公式": "f", "参考": "r"}.get(typ, "l")
        put(ws, f"A{r}", code, "r", None, "center")
        put(ws, f"B{r}", name, "l", None, "left")
        put(ws, f"C{r}", typ, typ_kind, None, "center")
        put(ws, f"D{r}", formula_txt, "r", None, "left", wrap=True)
        put(ws, f"E{r}", note, "l", None, "left", wrap=True)
        ws[f"D{r}"].font = Font(name=FONT, size=9, color=INK2)
        ws[f"E{r}"].font = Font(name=FONT, size=9, color=INK2)
        ws.row_dimensions[r].height = 20
        r += 1
    r += 1

    r = section(ws, r, "四、使用步骤", last_col)
    for i, s in enumerate(steps, 1):
        put(ws, f"A{r}", i, "s", None, "center")
        ws[f"A{r}"].fill = fill(NAVY2)
        ws[f"A{r}"].font = Font(name=FONT, size=10, bold=True, color=WHITE)
        ws.merge_cells(f"B{r}:{last_col}{r}")
        put(ws, f"B{r}", s, "l", None, "left", wrap=True)
        ws.row_dimensions[r].height = 20
        r += 1
    r += 1

    r = section(ws, r, "五、关联《收益管理操作法典》", last_col)
    for line in codex:
        r = para(ws, r, line, last_col, height=18)

    if extra_sections:
        for heading, lines in extra_sections:
            r += 1
            r = section(ws, r, heading, last_col)
            for line in lines:
                r = para(ws, r, line, last_col, height=18)

    r += 1
    r = para(ws, r, "杭州滨江云璟酒店 · 320 间房 · 中高端全服务 ｜ 酒店收益管理学院 HRMA ｜ 口径以 SPEC.md 为准",
             last_col, height=18, kind="n")
    return ws


# ---------------------------------------------------------------- 常用片段
WEEKDAY_CHOOSE = 'CHOOSE(WEEKDAY({c}{r},2),"周一","周二","周三","周四","周五","周六","周日")'

DEMAND_LEVELS = ["D1", "D2", "D3", "D4"]
TABLE_CODES = ["表A", "表B", "表C", "表D", "表E", "表G", "表H", "表L"]
ACTION_TYPES = ["调价", "MinLOS", "CTA", "关房", "关渠道", "超售", "升级"]
REVIEW_RESULT = ["做对", "做错", "中性", ""]

SPEC_MONTH_OCC = [0.58, 0.62, 0.74, 0.82, 0.84, 0.76, 0.70, 0.78, 0.83, 0.85, 0.78, 0.68]
SPEC_MONTH_ADR = [540, 560, 600, 660, 690, 620, 580, 630, 670, 700, 640, 570]

# SPEC 2.1 房型清单：(代码, 名称, 间数, 面积, 床型, 平峰BAR, 系数)
ROOM_TYPES = [
    ("STD-K", "标准大床房", 70, 32, "1.8m 大床", 580, 85),
    ("STD-T", "标准双床房", 50, 32, "1.2m×2", 580, 85),
    ("SUP-K", "高级大床房", 60, 38, "1.8m 大床", 680, 100),
    ("SUP-T", "高级双床房", 40, 38, "1.2m×2", 680, 100),
    ("DLX-K", "豪华大床房", 45, 45, "2.0m 大床", 820, 120),
    ("DLX-V", "豪华江景房", 15, 48, "2.0m 大床", 960, 140),
    ("EXE-K", "行政大床房", 25, 45, "2.0m 大床", 990, 145),
    ("EXE-T", "行政双床房", 5, 45, "1.2m×2", 990, 145),
    ("STE-B", "商务套房", 8, 65, "2.0m 大床", 1280, 180),
    ("STE-E", "行政套房", 2, 88, "2.0m 大床", 1880, 265),
]

# SPEC 2.4 竞争组合
COMPSET = [
    ("A", "江畔万枫酒店", "中高端", 280, 650, 1.2),
    ("B", "滨江希尔顿花园", "中高端", 240, 690, 2.5),
    ("C", "钱塘智选假日", "中端偏上", 300, 560, 3.1),
    ("D", "星光开元名都", "高端", 350, 880, 4.0),
    ("E", "云栖亚朵S", "中高端", 200, 620, 2.0),
]

# SPEC 2.3 客源细分
SEGMENTS = [
    ("FIT-WI", "上门散客", 0.08, 720, "DIR", 0.03),
    ("FIT-DIR", "官网/小程序直销", 0.17, 650, "DIR", 0.03),
    ("FIT-OTA", "OTA 散客", 0.31, 598, "OTA", 0.16),
    ("CORP", "协议公司客户", 0.22, 520, "CORP", 0.08),
    ("LOY", "会员（忠诚计划）", 0.12, 580, "LOY", 0.04),
    ("GRP", "团队（旅游/会议）", 0.08, 450, "WHO", 0.26),
    ("LNG", "长住客", 0.02, 380, "DIR", 0.03),
]

# 《法典》表 I · 预订窗口分档监控表
WINDOW_TABLE = [
    ("0-3天", "每日2次(10:00/17:00)", "OCC<70%", "上「今日特惠」尾房产品", None),
    ("4-7天", "每日1次", "落后10%", "启动48h闪购", 0.10),
    ("8-14天", "每日1次", "落后15%", "执行表A调价", 0.15),
    ("15-30天", "每周2次", "落后20%", "策略复盘+销售动员", 0.20),
    ("31-60天", "每周1次", "落后25%", "启动协议客户拜访计划", 0.25),
    (">60天", "每月1次", "落后30%", "上报L3，调整年度策略", 0.30),
]

# 《法典》表 K · 会议规模档
MEETING_SCALE = [
    ("小型", "≤30人", "≤15间夜", 8000, 1.10),
    ("中型", "31-120人", "16-60间夜", 30000, 1.00),
    ("大型", "121-300人", "61-150间夜", 80000, 0.95),
    ("超大型", ">300人", ">150间夜", 150000, 0.90),
]


def window_bucket_expr(cellref):
    """表 I 窗口分档（嵌套 IF，兼容 Excel 2016+/WPS）。"""
    c = cellref
    return (f'IF({c}="","",IF({c}<=3,"0-3天",IF({c}<=7,"4-7天",IF({c}<=14,"8-14天",'
            f'IF({c}<=30,"15-30天",IF({c}<=60,"31-60天",">60天"))))))')


def write_block(ws, r0, start_col, headers, rows, kinds=None, fmts=None, height=18):
    """写一个小参数块（如判定表），返回结束行+1。"""
    r = header_row(ws, r0, start_col, headers, height=24)
    for i, rw in enumerate(rows):
        for j, v in enumerate(rw):
            kind = "r"
            if kinds:
                kind = kinds[j] if isinstance(kinds[j], str) else kinds[j](i)
            fmt = fmts[j] if fmts else None
            put(ws, f"{get_column_letter(start_col + j)}{r}", v, kind, fmt, "left", wrap=True)
        ws.row_dimensions[r].height = height
        r += 1
    return r


def freeze(ws, ref):
    ws.freeze_panes = ref
