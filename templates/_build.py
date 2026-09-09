# -*- coding: utf-8 -*-
"""
HRMA · Excel 实操模板包 构建脚本

用法：
    /Users/abc/.workbuddy/binaries/python/envs/default/bin/python \
        /Users/abc/WorkBuddy/2026-09-03-08-48-09/hrma/templates/_build.py

产出 8 个工作簿到 templates/ 目录。所有计算单元格均为真实 Excel 公式。
口径来源：hrma/SPEC.md（唯一真相来源）
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import load_workbook  # noqa: E402

import _wb_part1 as p1  # noqa: E402
import _wb_part2 as p2  # noqa: E402

OUT = HERE

SPECS = [
    ("01_每日收益日报.xlsx", p1.build_01, ["说明", "日报", "数据字典"]),
    ("02_Pickup预订速度追踪表.xlsx", p1.build_02, ["说明", "Pickup追踪", "去年同期对照", "预测计算"]),
    ("03_竞争比价监测表.xlsx", p1.build_03, ["说明", "每日比价", "竞对档案", "ARI诊断"]),
    ("04_调价台账.xlsx", p1.build_04, ["说明", "台账", "命中率复盘", "观察清单"]),
    ("05_房型价差与价格体系.xlsx", p2.build_05, ["说明", "房型清单", "价格阶梯", "全年价格日历", "价差校验"]),
    ("06_年度预算与月度分解.xlsx", p2.build_06, ["说明", "年度假设", "月度预算", "分日预算", "滚动预测", "差异分析"]),
    ("07_周报月报模板.xlsx", p2.build_07, ["说明", "周报", "月报", "业主一页纸", "周末作战表"]),
    ("08_超额预订计算器.xlsx", p2.build_08, ["说明", "超售计算", "历史取消率", "团队询价决策", "会议综合贡献"]),
]


def check(path, expect_sheets):
    wb = load_workbook(path)
    names = wb.sheetnames
    errs = []
    if names != expect_sheets:
        errs.append(f"sheet 顺序不符：{names} != {expect_sheets}")
    if names[0] != "说明":
        errs.append("首张 sheet 不是『说明』")
    total = 0
    formula_cells = 0
    for ws in wb.worksheets:
        for rw in ws.iter_rows():
            for c in rw:
                if c.value is not None:
                    total += 1
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formula_cells += 1
    return names, total, formula_cells, errs


def main():
    print("=" * 78)
    print("HRMA Excel 实操模板包 · 构建")
    print("=" * 78)
    bad = 0
    for fname, builder, sheets in SPECS:
        path = os.path.join(OUT, fname)
        builder(path)
        names, total, nf, errs = check(path, sheets)
        size = os.path.getsize(path)
        status = "OK " if not errs else "ERR"
        print(f"[{status}] {fname:34s} sheets={len(names)}  单元格={total:6d}  公式={nf:6d}  {size/1024:7.1f} KB")
        for e in errs:
            bad += 1
            print(f"        !! {e}")
    print("=" * 78)
    if bad:
        print(f"发现 {bad} 处问题")
        return 1
    print("全部 8 个工作簿构建完成，sheet 结构与『说明』首页校验通过。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
